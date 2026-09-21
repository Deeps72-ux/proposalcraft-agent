import io
from datetime import datetime
from typing import Dict, Any

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, hex_color: str):
    """Set background shading color of a docx table cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color.replace('#', ''))
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def render_docx(proposal: Dict[str, Any]) -> bytes:
    """Generate a clean, structured Microsoft Word DOCX proposal using python-docx."""
    doc = Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette Constants
    COLOR_PRIMARY = RGBColor(15, 23, 42)     # Slate 900
    COLOR_ACCENT = RGBColor(37, 99, 235)     # Blue 600
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500

    client = proposal.get("client_name") or "Enterprise Client"
    project = proposal.get("project_title") or "Commercial Proposal & Technical Solution"
    company = proposal.get("company_name") or "ProposalCraft Solutions"
    created_at = proposal.get("created_at") or datetime.now().strftime("%B %d, %Y")
    sections = proposal.get("sections", {})
    review = proposal.get("review_feedback", {})

    # ==========================================
    # 1. COVER PAGE / HEADER
    # ==========================================
    p_tag = doc.add_paragraph()
    r_tag = p_tag.add_run("ENTERPRISE COMMERCIAL PROPOSAL")
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.font.bold = True
    r_tag.font.color.rgb = COLOR_ACCENT
    p_tag.paragraph_format.space_after = Pt(6)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(project)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(8)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(f"Prepared for {client}")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_MUTED
    p_sub.paragraph_format.space_after = Pt(24)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_items = [
        ("Submitted By:", company),
        ("Client Organization:", client),
        ("Submission Date:", str(created_at)[:10]),
        ("Quality Assurance:", f"Approved ({review.get('compliance_score', 95)}/100 Compliance)"),
    ]
    for i, (k, v) in enumerate(meta_items):
        r = meta_table.rows[i]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_margins(c1, 100, 100, 140, 140)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_PRIMARY

    doc.add_page_break()

    # Helper for adding Section Headings
    def add_section_heading(title: str):
        h = doc.add_paragraph()
        run = h.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_ACCENT
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        return h

    def add_sub_heading(title: str):
        h = doc.add_paragraph()
        run = h.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        return h

    def add_body_text(text: str):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_PRIMARY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        return p

    # ==========================================
    # 2. EXECUTIVE SUMMARY
    # ==========================================
    exec_sec = sections.get("executive_summary", {})
    add_section_heading("1. Executive Summary")
    add_sub_heading("Strategic Challenge & Context")
    add_body_text(exec_sec.get("challenge", "Context and objectives."))

    add_sub_heading("Proposed Solution Overview")
    add_body_text(exec_sec.get("proposed_solution", "Solution architecture."))

    strat_val = exec_sec.get("strategic_value", "")
    if strat_val:
        add_sub_heading("Strategic Value Proposition & ROI")
        add_body_text(strat_val)

    add_sub_heading("Key Highlights")
    for hl in exec_sec.get("highlights", []):
        bp = doc.add_paragraph(style="List Bullet")
        run = bp.add_run(hl)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        bp.paragraph_format.space_after = Pt(3)

    # ==========================================
    # 3. TECHNICAL ARCHITECTURE
    # ==========================================
    tech_sec = sections.get("technical_architecture", {})
    add_section_heading("2. Proposed Technical Architecture")
    add_body_text(tech_sec.get("overview", "Architecture specifications."))

    comps = tech_sec.get("components", [])
    if comps:
        add_sub_heading("Core Components & Technology Stack")
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].width = Inches(1.8)
        hdr_cells[1].width = Inches(1.8)
        hdr_cells[2].width = Inches(2.9)

        headers = ["Component", "Technology Stack", "Description & Responsibility"]
        for idx, title in enumerate(headers):
            cell = hdr_cells[idx]
            set_cell_background(cell, "2563EB")
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(title)
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for c in comps:
            row_cells = table.add_row().cells
            row_cells[0].width = Inches(1.8)
            row_cells[1].width = Inches(1.8)
            row_cells[2].width = Inches(2.9)

            for i, text in enumerate([c.get("name", ""), c.get("tech_stack", ""), c.get("description", "")]):
                cell = row_cells[i]
                set_cell_margins(cell, 80, 80, 100, 100)
                p = cell.paragraphs[0]
                r = p.add_run(text)
                r.font.size = Pt(9)
                if i == 0:
                    r.font.bold = True

    add_sub_heading("Security, Privacy & Compliance")
    add_body_text(tech_sec.get("security_and_compliance", "Encryption & Governance."))

    add_sub_heading("Scalability & Resilience")
    add_body_text(tech_sec.get("scalability_and_resilience", "High-availability clustering."))

    # ==========================================
    # 4. SCOPE & DELIVERABLES
    # ==========================================
    deliv_sec = sections.get("deliverables", {})
    add_section_heading("3. Scope & Key Deliverables")

    phases = deliv_sec.get("phases", [])
    if phases:
        deliv_table = doc.add_table(rows=1, cols=3)
        deliv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        d_hdr = deliv_table.rows[0].cells
        d_hdr[0].width = Inches(1.8)
        d_hdr[1].width = Inches(2.9)
        d_hdr[2].width = Inches(1.8)

        for idx, title in enumerate(["Phase & Duration", "Key Deliverables", "Strategic Outcome"]):
            cell = d_hdr[idx]
            set_cell_background(cell, "0F172A")
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(title)
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for p_item in phases:
            row_cells = deliv_table.add_row().cells
            row_cells[0].width = Inches(1.8)
            row_cells[1].width = Inches(2.9)
            row_cells[2].width = Inches(1.8)

            # Cell 0: Phase
            set_cell_margins(row_cells[0], 80, 80, 100, 100)
            p0 = row_cells[0].paragraphs[0]
            r0 = p0.add_run(f"{p_item.get('phase', '')}\n({p_item.get('duration', '')})")
            r0.font.bold = True
            r0.font.size = Pt(9)

            # Cell 1: Deliverables
            set_cell_margins(row_cells[1], 80, 80, 100, 100)
            p1 = row_cells[1].paragraphs[0]
            d_list = "\n".join([f"• {d}" for d in p_item.get("deliverables", [])])
            r1 = p1.add_run(d_list)
            r1.font.size = Pt(9)

            # Cell 2: Outcome
            set_cell_margins(row_cells[2], 80, 80, 100, 100)
            p2 = row_cells[2].paragraphs[0]
            r2 = p2.add_run(p_item.get("outcome", ""))
            r2.font.size = Pt(9)

    # ==========================================
    # 5. COMMERCIAL & PRICING
    # ==========================================
    price_sec = sections.get("pricing", {})
    currency = price_sec.get("currency", "USD")
    total_inv = price_sec.get("total_investment", 0.0)

    add_section_heading("4. Commercial & Milestone Investment")

    milestones = price_sec.get("milestones", [])
    if milestones:
        ptable = doc.add_table(rows=1, cols=5)
        ptable.alignment = WD_TABLE_ALIGNMENT.CENTER
        p_hdr = ptable.rows[0].cells
        col_widths = [Inches(1.8), Inches(2.2), Inches(0.9), Inches(1.1), Inches(0.5)]

        for idx, title in enumerate(["Milestone", "Trigger Deliverable", "Timeline", f"Amount ({currency})", "Share"]):
            cell = p_hdr[idx]
            cell.width = col_widths[idx]
            set_cell_background(cell, "1E3A8A")
            set_cell_margins(cell, 100, 100, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(title)
            r.font.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

        for m in milestones:
            row_cells = ptable.add_row().cells
            for idx, w in enumerate(col_widths):
                row_cells[idx].width = w
                set_cell_margins(row_cells[idx], 80, 80, 80, 80)

            amt = float(m.get("amount", 0))
            row_cells[0].paragraphs[0].add_run(m.get("milestone", "")).font.size = Pt(9)
            row_cells[1].paragraphs[0].add_run(m.get("deliverable", "")).font.size = Pt(9)
            row_cells[2].paragraphs[0].add_run(m.get("timeline", "")).font.size = Pt(9)
            row_cells[3].paragraphs[0].add_run(f"{amt:,.2f}").font.size = Pt(9)
            row_cells[4].paragraphs[0].add_run(str(m.get("percentage", ""))).font.size = Pt(9)

        # Summary row
        tot_cells = ptable.add_row().cells
        for idx, w in enumerate(col_widths):
            tot_cells[idx].width = w
            set_cell_background(tot_cells[idx], "FEF3C7")
            set_cell_margins(tot_cells[idx], 90, 90, 80, 80)

        tot_cells[0].paragraphs[0].add_run("Total Contract Investment").font.bold = True
        tot_cells[3].paragraphs[0].add_run(f"{total_inv:,.2f} {currency}").font.bold = True
        tot_cells[4].paragraphs[0].add_run("100%").font.bold = True

    terms = price_sec.get("payment_terms", "")
    if terms:
        add_sub_heading("Payment Terms")
        add_body_text(terms)

    assumptions = price_sec.get("assumptions", [])
    if assumptions:
        add_sub_heading("Commercial Assumptions")
        for a in assumptions:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(a).font.size = Pt(9.5)
            bp.paragraph_format.space_after = Pt(2)

    # Save to binary buffer
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
