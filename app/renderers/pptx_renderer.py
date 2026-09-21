import io
from datetime import datetime
from typing import Dict, Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def apply_solid_fill(shape, rgb: RGBColor):
    """Fill a shape with a solid RGB color."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()  # No border by default


def set_slide_background(slide, prs, rgb: RGBColor):
    """Add a full-slide rectangle background with specified RGB color."""
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    apply_solid_fill(bg, rgb)
    # Send to back if possible, or create it first before adding other elements


def create_header(slide, title_text: str, category_text: str = "PROPOSALCRAFT EXECUTIVE BRIEF"):
    """Render consistent executive slide header."""
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.name = "Arial"
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = RGBColor(59, 130, 246)  # Blue 500
    p_cat.space_after = Pt(2)

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900


def render_pptx(proposal: Dict[str, Any]) -> bytes:
    """Generate a 5-slide widescreen executive pitch deck using python-pptx."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]  # Blank layout

    # Color Palette
    DARK_NAVY = RGBColor(15, 23, 42)      # Slate 900
    ACCENT_BLUE = RGBColor(37, 99, 235)   # Blue 600
    LIGHT_BG = RGBColor(248, 250, 252)    # Slate 50
    CARD_BG = RGBColor(255, 255, 255)     # White
    TEXT_MAIN = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200

    client = proposal.get("client_name") or "Enterprise Client"
    project = proposal.get("project_title") or "Commercial & Technical Proposal"
    company = proposal.get("company_name") or "ProposalCraft Solutions"
    created_at = proposal.get("created_at") or datetime.now().strftime("%B %d, %Y")
    sections = proposal.get("sections", {})
    review = proposal.get("review_feedback", {})

    # =========================================================================
    # SLIDE 1: TITLE / COVER SLIDE (Dark Modern Enterprise Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1, prs, DARK_NAVY)

    # Accent decorative bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(0.15), Inches(3.2))
    apply_solid_fill(bar, ACCENT_BLUE)

    title_box = slide1.shapes.add_textbox(Inches(1.4), Inches(1.7), Inches(10.5), Inches(3.4))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "EXECUTIVE PITCH DECK & COMMERCIAL PROPOSAL"
    p_badge.font.name = "Arial"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(96, 165, 250)
    p_badge.space_after = Pt(10)

    p_main = tf1.add_paragraph()
    p_main.text = project
    p_main.font.name = "Arial"
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = RGBColor(255, 255, 255)
    p_main.space_after = Pt(12)

    p_for = tf1.add_paragraph()
    p_for.text = f"Prepared Exclusively for: {client}"
    p_for.font.name = "Arial"
    p_for.font.size = Pt(16)
    p_for.font.color.rgb = RGBColor(203, 213, 225)

    # Footer Card on Slide 1
    meta_box = slide1.shapes.add_textbox(Inches(1.4), Inches(5.8), Inches(10.5), Inches(1.0))
    tf_meta = meta_box.text_frame
    p_meta = tf_meta.paragraphs[0]
    p_meta.text = f"Presented by {company}   |   Date: {str(created_at)[:10]}   |   Compliance Rating: {review.get('compliance_score', 95)}/100"
    p_meta.font.name = "Arial"
    p_meta.font.size = Pt(11)
    p_meta.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: EXECUTIVE SUMMARY & STRATEGIC VALUE
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2, prs, LIGHT_BG)
    create_header(slide2, "Executive Summary & Strategic Value Proposition")

    exec_sec = sections.get("executive_summary", {})
    card_width = Inches(3.64)
    card_height = Inches(4.8)
    card_top = Inches(1.8)

    col_data = [
        ("The Strategic Challenge", exec_sec.get("challenge", "Context and client problem statement."), RGBColor(239, 68, 68)),
        ("The Proposed Solution", exec_sec.get("proposed_solution", "Tailored engineering solution."), ACCENT_BLUE),
        ("Strategic Value & ROI", exec_sec.get("strategic_value", "Measurable enterprise return on investment."), RGBColor(16, 185, 129)),
    ]

    for idx, (head, content, badge_color) in enumerate(col_data):
        left = Inches(0.8 + idx * 4.0)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, card_top, card_width, card_height)
        apply_solid_fill(card, CARD_BG)
        card.line.color.rgb = CARD_BORDER

        tb = slide2.shapes.add_textbox(left + Inches(0.3), card_top + Inches(0.3), card_width - Inches(0.6), card_height - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = head
        p_h.font.name = "Arial"
        p_h.font.size = Pt(15)
        p_h.font.bold = True
        p_h.font.color.rgb = badge_color
        p_h.space_after = Pt(12)

        p_body = tf.add_paragraph()
        p_body.text = content
        p_body.font.name = "Arial"
        p_body.font.size = Pt(10.5)
        p_body.font.color.rgb = TEXT_MAIN
        p_body.space_after = Pt(14)

        if idx == 2 and exec_sec.get("highlights"):
            p_hl_title = tf.add_paragraph()
            p_hl_title.text = "Key Takeaways:"
            p_hl_title.font.bold = True
            p_hl_title.font.size = Pt(10)
            p_hl_title.font.color.rgb = TEXT_MAIN
            p_hl_title.space_after = Pt(4)
            for hl in exec_sec.get("highlights", [])[:3]:
                p_item = tf.add_paragraph()
                p_item.text = f"• {hl}"
                p_item.font.size = Pt(9.5)
                p_item.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: PROPOSED TECHNICAL ARCHITECTURE
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3, prs, LIGHT_BG)
    create_header(slide3, "Proposed Technical Architecture & Tech Stack")

    tech_sec = sections.get("technical_architecture", {})
    comps = tech_sec.get("components", [])[:4]

    # Overview Banner
    banner = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.9))
    apply_solid_fill(banner, RGBColor(238, 242, 255))
    banner.line.color.rgb = RGBColor(199, 210, 254)

    tb_b = slide3.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.3), Inches(0.8))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p_bo = tf_b.paragraphs[0]
    p_bo.text = tech_sec.get("overview", "Modular, scalable enterprise architecture designed for resilience and performance.")
    p_bo.font.name = "Arial"
    p_bo.font.size = Pt(10)
    p_bo.font.color.rgb = DARK_NAVY

    # 4 Components grid
    c_width = Inches(2.7)
    c_height = Inches(3.2)
    for i, comp in enumerate(comps):
        c_left = Inches(0.8 + i * 3.0)
        c_top = Inches(3.0)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, c_width, c_height)
        apply_solid_fill(card, CARD_BG)
        card.line.color.rgb = CARD_BORDER

        ctb = slide3.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), c_width - Inches(0.4), c_height - Inches(0.4))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = comp.get("name", "Module")
        cp0.font.name = "Arial"
        cp0.font.size = Pt(13)
        cp0.font.bold = True
        cp0.font.color.rgb = ACCENT_BLUE
        cp0.space_after = Pt(8)

        cp1 = ctf.add_paragraph()
        cp1.text = "Tech Stack:"
        cp1.font.size = Pt(9.5)
        cp1.font.bold = True
        cp1.font.color.rgb = TEXT_MUTED

        cp2 = ctf.add_paragraph()
        cp2.text = comp.get("tech_stack", "")
        cp2.font.size = Pt(10)
        cp2.font.bold = True
        cp2.font.color.rgb = TEXT_MAIN
        cp2.space_after = Pt(8)

        cp3 = ctf.add_paragraph()
        cp3.text = comp.get("description", "")
        cp3.font.size = Pt(9)
        cp3.font.color.rgb = TEXT_MUTED

    # Bottom Governance bar
    gov_box = slide3.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6))
    gtf = gov_box.text_frame
    gtf.word_wrap = True
    gp = gtf.paragraphs[0]
    gp.text = f"Security & Scalability: {tech_sec.get('security_and_compliance', '')[:140]}..."
    gp.font.name = "Arial"
    gp.font.size = Pt(9)
    gp.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: SCOPE & PHASED IMPLEMENTATION ROADMAP
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4, prs, LIGHT_BG)
    create_header(slide4, "Scope of Work & Phased Implementation Roadmap")

    deliv_sec = sections.get("deliverables", {})
    phases = deliv_sec.get("phases", [])[:4]

    p_width = Inches(2.7)
    p_height = Inches(4.7)

    for idx, p_item in enumerate(phases):
        p_left = Inches(0.8 + idx * 3.0)
        p_top = Inches(1.9)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, p_left, p_top, p_width, p_height)
        apply_solid_fill(card, CARD_BG)
        card.line.color.rgb = CARD_BORDER

        # Phase Header Tag
        tag = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, p_left, p_top, p_width, Inches(0.7))
        apply_solid_fill(tag, DARK_NAVY if idx % 2 == 0 else ACCENT_BLUE)

        tb_tag = slide4.shapes.add_textbox(p_left + Inches(0.15), p_top + Inches(0.08), p_width - Inches(0.3), Inches(0.55))
        tf_tag = tb_tag.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = p_item.get("phase", f"Phase {idx+1}")
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(10.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = RGBColor(255, 255, 255)

        p_subtag = tf_tag.add_paragraph()
        p_subtag.text = p_item.get("duration", "")
        p_subtag.font.size = Pt(8.5)
        p_subtag.font.color.rgb = RGBColor(226, 232, 240)

        # Deliverables content
        tb_body = slide4.shapes.add_textbox(p_left + Inches(0.2), p_top + Inches(0.85), p_width - Inches(0.4), p_height - Inches(1.0))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True

        p_dt = tf_body.paragraphs[0]
        p_dt.text = "Key Deliverables:"
        p_dt.font.size = Pt(9.5)
        p_dt.font.bold = True
        p_dt.font.color.rgb = TEXT_MAIN
        p_dt.space_after = Pt(4)

        for d in p_item.get("deliverables", [])[:4]:
            p_d = tf_body.add_paragraph()
            p_d.text = f"• {d}"
            p_d.font.size = Pt(8.5)
            p_d.font.color.rgb = TEXT_MUTED
            p_d.space_after = Pt(2)

        if p_item.get("outcome"):
            p_oc_title = tf_body.add_paragraph()
            p_oc_title.text = "Milestone Outcome:"
            p_oc_title.font.size = Pt(9.5)
            p_oc_title.font.bold = True
            p_oc_title.font.color.rgb = ACCENT_BLUE
            p_oc_title.space_before = Pt(8)
            p_oc_title.space_after = Pt(2)

            p_oc = tf_body.add_paragraph()
            p_oc.text = p_item.get("outcome", "")
            p_oc.font.size = Pt(8.5)
            p_oc.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 5: COMMERCIAL INVESTMENT & NEXT STEPS
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5, prs, LIGHT_BG)
    create_header(slide5, "Commercial Investment & Milestone Schedule")

    price_sec = sections.get("pricing", {})
    currency = price_sec.get("currency", "USD")
    total_inv = price_sec.get("total_investment", 0.0)
    milestones = price_sec.get("milestones", [])

    # Left: Milestone Table Card
    table_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
    apply_solid_fill(table_card, CARD_BG)
    table_card.line.color.rgb = CARD_BORDER

    tb_t = slide5.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(6.9), Inches(4.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True

    p_t_h = tf_t.paragraphs[0]
    p_t_h.text = "Milestone Payment Schedule"
    p_t_h.font.name = "Arial"
    p_t_h.font.size = Pt(14)
    p_t_h.font.bold = True
    p_t_h.font.color.rgb = DARK_NAVY
    p_t_h.space_after = Pt(10)

    for m in milestones[:4]:
        pm = tf_t.add_paragraph()
        pm.text = f"{m.get('milestone', '')} ({m.get('timeline', '')})"
        pm.font.size = Pt(10)
        pm.font.bold = True
        pm.font.color.rgb = ACCENT_BLUE

        pm_sub = tf_t.add_paragraph()
        amt = float(m.get('amount', 0))
        pm_sub.text = f"Deliverable: {m.get('deliverable', '')} | Amount: {amt:,.2f} {currency} ({m.get('percentage', '')})"
        pm_sub.font.size = Pt(9)
        pm_sub.font.color.rgb = TEXT_MUTED
        pm_sub.space_after = Pt(6)

    # Right Column: Total Investment & Next Steps Card
    right_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.8), Inches(3.9), Inches(4.8))
    apply_solid_fill(right_card, DARK_NAVY)

    tb_r = slide5.shapes.add_textbox(Inches(8.9), Inches(2.1), Inches(3.3), Inches(4.2))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_tot_lbl = tf_r.paragraphs[0]
    p_tot_lbl.text = "TOTAL INVESTMENT"
    p_tot_lbl.font.name = "Arial"
    p_tot_lbl.font.size = Pt(11)
    p_tot_lbl.font.bold = True
    p_tot_lbl.font.color.rgb = RGBColor(148, 163, 184)
    p_tot_lbl.space_after = Pt(2)

    p_tot_val = tf_r.add_paragraph()
    p_tot_val.text = f"{total_inv:,.2f} {currency}"
    p_tot_val.font.name = "Arial"
    p_tot_val.font.size = Pt(26)
    p_tot_val.font.bold = True
    p_tot_val.font.color.rgb = RGBColor(254, 240, 138)  # Light Amber
    p_tot_val.space_after = Pt(14)

    p_terms_h = tf_r.add_paragraph()
    p_terms_h.text = "Commercial Terms:"
    p_terms_h.font.size = Pt(10)
    p_terms_h.font.bold = True
    p_terms_h.font.color.rgb = RGBColor(255, 255, 255)

    p_terms = tf_r.add_paragraph()
    p_terms.text = price_sec.get("payment_terms", "Net 30 days upon milestone acceptance.")
    p_terms.font.size = Pt(8.5)
    p_terms.font.color.rgb = RGBColor(203, 213, 225)
    p_terms.space_after = Pt(12)

    p_next_h = tf_r.add_paragraph()
    p_next_h.text = "Immediate Next Steps:"
    p_next_h.font.size = Pt(10)
    p_next_h.font.bold = True
    p_next_h.font.color.rgb = RGBColor(255, 255, 255)

    steps = [
        "1. Executive alignment & contract execution",
        "2. Technical kickoff & architecture sprint",
        "3. Initial milestone sprint delivery"
    ]
    for s in steps:
        ps = tf_r.add_paragraph()
        ps.text = s
        ps.font.size = Pt(8.5)
        ps.font.color.rgb = RGBColor(203, 213, 225)
        ps.space_after = Pt(2)

    # Save to binary buffer
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
