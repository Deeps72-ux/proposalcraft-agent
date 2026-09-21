import io
from pptx import Presentation
from docx import Document
from app.renderers.pdf_renderer import render_pdf
from app.renderers.docx_renderer import render_docx
from app.renderers.pptx_renderer import render_pptx


SAMPLE_PROPOSAL = {
    "id": "test-proposal-123",
    "client_name": "Apex Global Solutions",
    "project_title": "Enterprise Cloud & AI Modernization Proposal",
    "company_name": "ProposalCraft Solutions",
    "currency": "USD",
    "created_at": "2026-09-21T12:00:00Z",
    "sections": {
        "executive_summary": {
            "title": "1. Executive Summary",
            "challenge": "Client requires an upgrade to cloud-native microservices with enterprise resilience.",
            "proposed_solution": "ProposalCraft delivers automated containerized orchestration and scalable APIs.",
            "strategic_value": "Reduces operational overhead by 40% and achieves 99.99% availability.",
            "highlights": [
                "Zero-downtime migration strategy",
                "Automated CI/CD security pipelines",
                "Full SOC2 Type II compliance",
            ],
        },
        "technical_architecture": {
            "title": "2. Proposed Technical Architecture",
            "overview": "Modern event-driven architecture running on Kubernetes with distributed caching.",
            "components": [
                {
                    "name": "API Gateway",
                    "tech_stack": "FastAPI, Python 3.11",
                    "description": "High-throughput asynchronous ingestion layer.",
                },
                {
                    "name": "Agent Engine",
                    "tech_stack": "LangGraph, Groq LLM",
                    "description": "Autonomous workflow state machine.",
                },
            ],
            "security_and_compliance": "AES-256 at-rest, TLS 1.3 in-transit, role-based access control.",
            "scalability_and_resilience": "Multi-region autoscaling with automated failover.",
        },
        "deliverables": {
            "title": "3. Scope & Key Deliverables",
            "phases": [
                {
                    "phase": "Phase 1: Architecture & Blueprint",
                    "duration": "Weeks 1-3",
                    "deliverables": ["System Design Document", "Security Baseline Review"],
                    "outcome": "Approved architecture specification.",
                },
                {
                    "phase": "Phase 2: Core Engineering & Testing",
                    "duration": "Weeks 4-8",
                    "deliverables": ["Core Platform Implementation", "Automated Test Suite"],
                    "outcome": "Production-ready MVP.",
                },
            ],
        },
        "pricing": {
            "title": "4. Commercial & Milestone Investment",
            "currency": "USD",
            "milestones": [
                {
                    "milestone": "Milestone 1: Design Sign-off",
                    "deliverable": "Approved Architecture Spec",
                    "timeline": "Week 3",
                    "amount": 25000.0,
                    "percentage": "25%",
                },
                {
                    "milestone": "Milestone 2: Final Acceptance",
                    "deliverable": "Production Deployment & Handover",
                    "timeline": "Week 8",
                    "amount": 75000.0,
                    "percentage": "75%",
                },
            ],
            "total_investment": 100000.0,
            "payment_terms": "Net 30 days upon milestone sign-off.",
            "assumptions": ["Cloud infrastructure hosted on client AWS account."],
        },
    },
    "review_feedback": {
        "compliance_score": 98,
        "tone_score": 96,
        "executive_ready": True,
        "strengths": ["Clear technical alignment", "Predictable phased commercial schedule"],
        "recommendations": ["Reiterate SLA parameters during contract review"],
        "review_notes": "Proposal meets all enterprise standards and is cleared for distribution.",
    },
}


def test_pdf_renderer():
    """Verify ReportLab PDF generation creates valid binary output."""
    pdf_bytes = render_pdf(SAMPLE_PROPOSAL)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 2000
    assert pdf_bytes.startswith(b"%PDF"), "PDF must start with %PDF magic header"


def test_docx_renderer():
    """Verify python-docx generation creates valid Word document."""
    docx_bytes = render_docx(SAMPLE_PROPOSAL)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 5000
    # DOCX is a ZIP archive starting with PK
    assert docx_bytes.startswith(b"PK\x03\x04")

    # Verify readable by python-docx
    doc = Document(io.BytesIO(docx_bytes))
    paragraphs_text = " ".join([p.text for p in doc.paragraphs])
    assert "Apex Global Solutions" in paragraphs_text
    assert "Executive Summary" in paragraphs_text


def test_pptx_renderer():
    """Verify python-pptx generation creates valid 5-slide pitch presentation."""
    pptx_bytes = render_pptx(SAMPLE_PROPOSAL)
    assert isinstance(pptx_bytes, bytes)
    assert len(pptx_bytes) > 5000
    assert pptx_bytes.startswith(b"PK\x03\x04")

    # Verify readable by python-pptx and contains exactly 5 slides
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 5, f"Expected 5 slides, got {len(prs.slides)}"
